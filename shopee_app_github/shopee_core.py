"""Core logic for the Shopee Bulk Stock Update app — stateful reconciler (no Streamlit dep)."""
import io, re, zipfile
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ACC_PRESET = 10

# ---------- helpers ----------
def fix_shopee_bytes(data: bytes) -> io.BytesIO:
    zin = zipfile.ZipFile(io.BytesIO(data), "r")
    out = io.BytesIO(); zout = zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED)
    for item in zin.infolist():
        raw = zin.read(item.filename)
        if item.filename.endswith(".xml"):
            raw = re.sub(r'activePane="[^"]*"', 'activePane="topLeft"', raw.decode("utf-8","ignore")).encode("utf-8")
        zout.writestr(item, raw)
    zin.close(); zout.close(); out.seek(0); return out

def _nid(v):
    if v is None: return ""
    s = str(v).strip()
    return s[:-2] if (s.endswith(".0") and s[:-2].isdigit()) else s

def _split(cell):
    return [x for x in re.split(r"[;,/]", str(cell))] if cell is not None else []

def is_freebie(model): return "FREEB" in (model or "").upper()

def is_skip_sku(sku):
    s = re.sub(r"\s+", " ", str(sku or "").strip().upper())
    return s == "NO" or bool(re.search(r"NO\s*/\s*0\s*QTY", s))

ACC_KW = ["tempered glass","screen protector","protection plan","warranty","coverage","adapter",
          "charger","cable","earpiece","dummy","hypercharge","power bank","powerbank","protector",
          "full coverage","auto fit","usb-c adapter","fast charging"]
def is_accessory(pname):
    seg = str(pname or "").lower().split("|")[0]
    return any(k in seg for k in ACC_KW)

def _find_col(ws, hr, *kw):
    vals = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[hr]]
    for k in kw:
        for i, v in enumerate(vals):
            if k in v: return i + 1
    return None

# ---------- Masterlist ----------
def load_masterlist_full(file):
    wb = openpyxl.load_workbook(file, data_only=True); ws = wb.active
    hr = idc = qc = cc = bc = mc = colc = None
    for r in range(1, min(6, ws.max_row)+1):
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[r]]
        if "stock type id" in vals:
            hr = r; idc = vals.index("stock type id")+1
            cc = vals.index("category")+1 if "category" in vals else None
            bc = vals.index("brand")+1 if "brand" in vals else None
            mc = vals.index("model")+1 if "model" in vals else None
            colc = vals.index("color")+1 if "color" in vals else None
            qc = vals.index("total")+1 if "total" in vals else None
            break
    if not idc: raise ValueError("Masterlist missing 'Stock Type ID' column.")
    if not qc: raise ValueError("Masterlist missing 'Total' (Available Quantity) column.")
    rows = []
    for r in range(hr+1, ws.max_row+1):
        sid = _nid(ws.cell(r, idc).value)
        if not sid: continue
        q = ws.cell(r, qc).value
        try: q = int(float(q))
        except (TypeError, ValueError): q = 0
        rows.append({"id":sid,"category":(ws.cell(r,cc).value if cc else "") or "",
                     "brand":(ws.cell(r,bc).value if bc else "") or "",
                     "model":(ws.cell(r,mc).value if mc else "") or "",
                     "color":(ws.cell(r,colc).value if colc else "") or "","avail":q})
    return rows

# ---------- Shopee variations ----------
def read_shopee_variations(data):
    wb = openpyxl.load_workbook(fix_shopee_bytes(data), data_only=True); ws = wb.active
    vcol = scol = pidc = pnc = vnc = skuc = hr = None
    for r in range(1, min(4, ws.max_row)+1):
        vals = [str(c.value).strip().lower() if c.value is not None else "" for c in ws[r]]
        for i,v in enumerate(vals):
            if v=="et_title_variation_id" or v=="variation id": vcol=i+1
            if v.startswith("et_title_variation_sto") or v=="seller stock": scol=i+1
            if v=="et_title_product_id" or v=="product id": pidc=i+1
            if v=="et_title_product_name" or v=="product name": pnc=i+1
            if v=="et_title_variation_name" or v=="variation name": vnc=i+1
            if v=="et_title_variation_sku" or v=="sku": skuc=i+1
        if vcol and scol: hr=r; break
    out=[]
    for r in range(hr+1, ws.max_row+1):
        vid=_nid(ws.cell(r,vcol).value)
        if not vid or not vid.isdigit(): continue
        out.append({"vid":vid,
            "pid":_nid(ws.cell(r,pidc).value) if pidc else "",
            "pname":(ws.cell(r,pnc).value if pnc else "") or "",
            "vname":(ws.cell(r,vnc).value if vnc else "") or "",
            "sku":(ws.cell(r,skuc).value if skuc else "") or "",
            "stock":ws.cell(r,scol).value})
    return out

# ---------- Registry: read locks + apply decisions ----------
def read_registry(reg_bytes):
    """Return locked{vid:[mids]}, not_selling set, not_on_shopee set, and a decisions summary,
    AFTER folding in decisions from 'New Masterlist SKUs' and 'Match Review' tabs."""
    wb = openpyxl.load_workbook(io.BytesIO(reg_bytes), data_only=True)
    names = {n.lower(): n for n in wb.sheetnames}
    def sh(*subs):
        for low,real in names.items():
            if all(s in low for s in subs): return wb[real]
        return None
    locked={}; not_selling=set(); not_on_shopee=set()
    summary={"linked":0,"not_selling":0,"not_on_shopee":0,"match_review_locked":0}
    def add_lock(vid,mid):
        if not vid or not mid: return
        locked.setdefault(vid,[])
        if mid not in locked[vid]: locked[vid].append(mid)
    ws=sh("locked")
    if ws is not None:
        vc=_find_col(ws,1,"variation id"); mc=_find_col(ws,1,"masterlist id")
        if vc and mc:
            for r in range(2,ws.max_row+1):
                vid=_nid(ws.cell(r,vc).value)
                for m in [_nid(x) for x in _split(ws.cell(r,mc).value)]:
                    add_lock(vid,m)
    for key,subs in [("ns",("not selling",)),("nos",("not on shopee",))]:
        ws=sh(*subs)
        if ws is not None:
            sc=_find_col(ws,1,"stock type id","masterlist id")
            if sc:
                for r in range(2,ws.max_row+1):
                    sid=_nid(ws.cell(r,sc).value)
                    if sid: (not_selling if key=="ns" else not_on_shopee).add(sid)
    # apply New Masterlist SKUs decisions
    ws=sh("new masterlist")
    if ws is not None:
        mc=_find_col(ws,1,"stock type id","masterlist id"); vc=_find_col(ws,1,"variation id"); dc=_find_col(ws,1,"decision")
        if mc and dc:
            for r in range(2,ws.max_row+1):
                mid=_nid(ws.cell(r,mc).value); dec=str(ws.cell(r,dc).value or "").strip().lower()
                vid=_nid(ws.cell(r,vc).value) if vc else ""
                if not mid or not dec: continue
                if dec.startswith("link") and vid: add_lock(vid,mid); summary["linked"]+=1
                elif "not applicable" in dec: not_selling.add(mid); summary["not_selling"]+=1
                elif "not on shopee" in dec: not_on_shopee.add(mid); summary["not_on_shopee"]+=1
    # apply Match Review decisions
    ws=sh("match review")
    if ws is not None:
        vc=_find_col(ws,1,"variation id"); dc=_find_col(ws,1,"decision")
        corr=_find_col(ws,1,"corrected masterlist id"); matched=_find_col(ws,1,"matched masterlist id"); altc=_find_col(ws,1,"alt")
        if vc and dc:
            for r in range(2,ws.max_row+1):
                vid=_nid(ws.cell(r,vc).value); dec=str(ws.cell(r,dc).value or "").strip().lower()
                if not vid or not dec: continue
                mid=""
                if "use alt" in dec and altc: mid=_nid(ws.cell(r,altc).value)
                elif "confirm" in dec and matched: mid=_nid(ws.cell(r,matched).value)
                if not mid and corr: mid=_nid(ws.cell(r,corr).value)
                if mid and (dec.startswith(("confirm","use alt","corrected","link"))):
                    add_lock(vid,mid); summary["match_review_locked"]+=1
    return locked, not_selling, not_on_shopee, summary

# ---------- Reconcile ----------
def reconcile(ml_full, locked, not_selling, not_on_shopee, shopee_vars):
    avail={m["id"]:m["avail"] for m in ml_full}
    mlmap={m["id"]:m for m in ml_full}
    # classify each Shopee variation
    cls={"locked":[],"skipped":[],"accessories":[],"match_review":[]}
    for sv in shopee_vars:
        vid=sv["vid"]
        if vid in locked:
            mids=locked[vid]; sv2=dict(sv); sv2["mids"]=mids
            sv2["target"]=sum(avail.get(m,0) for m in mids)
            sv2["ml_desc"]="; ".join(f"{m}:{mlmap.get(m,{}).get('model','?')}|{mlmap.get(m,{}).get('color','')}" for m in mids)
            cls["locked"].append(sv2)
        elif is_skip_sku(sv["sku"]):
            cls["skipped"].append(sv)
        elif is_accessory(sv["pname"]):
            cls["accessories"].append(sv)
        else:
            cls["match_review"].append(sv)
    covered=set(m for mids in locked.values() for m in mids)|not_selling|not_on_shopee
    new_ml=[m for m in ml_full if (m["id"] not in covered) and not is_freebie(m["model"])]
    return cls, new_ml, avail, mlmap

# ---------- Build registry workbook ----------
def _style(ws, headers, kinds):
    hf=PatternFill("solid",fgColor="1F3864"); hfont=Font(name="Arial",bold=True,color="FFFFFF",size=10)
    YEL=PatternFill("solid",fgColor="FFD966"); ORG=PatternFill("solid",fgColor="F4B183")
    hblack=Font(name="Arial",bold=True,color="000000",size=10)
    thin=Side(style="thin",color="D9D9D9"); bd=Border(left=thin,right=thin,top=thin,bottom=thin)
    for c in range(1,len(headers)+1):
        cell=ws.cell(1,c); cell.value=headers[c-1]
        cell.font=hfont; cell.fill=hf; cell.alignment=Alignment(horizontal="center",vertical="center",wrap_text=True); cell.border=bd
        if c-1 < len(kinds):
            if kinds[c-1]=="M": cell.fill=YEL; cell.font=hblack
            elif kinds[c-1]=="S": cell.fill=ORG; cell.font=hblack

def build_registry_workbook(cls, new_ml, mlmap, not_selling, not_on_shopee, summary):
    wb=openpyxl.Workbook(); base=Font(name="Arial",size=10)
    editf=PatternFill("solid",fgColor="FFF2CC")
    # Summary
    s=wb.active; s.title="Summary"; s.sheet_view.showGridLines=False
    s["B2"]="Shopee Stock Bulk Update — Match Registry"; s["B2"].font=Font(name="Arial",bold=True,size=15,color="1F3864")
    rows=[("Locked Matches (sync stock)",len(cls["locked"])),
          ("New Masterlist SKUs to review",len(new_ml)),
          ("Match Review (Shopee SKUs to match)",len(cls["match_review"])),
          ("Skipped (SKU=NO / 0 Qty)",len(cls["skipped"])),
          ("Accessories (preset 10)",len(cls["accessories"])),
          ("Not Selling in Shopee",len(not_selling)),
          ("Not on Shopee Yet",len(not_on_shopee))]
    s["B4"]="Category"; s["C4"]="Count"
    for c in ("B4","C4"): s[c].font=Font(name="Arial",bold=True,color="FFFFFF"); s[c].fill=PatternFill("solid",fgColor="1F3864")
    for i,(lab,cnt) in enumerate(rows,start=5):
        s.cell(i,2,lab).font=base; s.cell(i,3,cnt).font=base
    note=("Applied from your review: %d linked, %d not-selling, %d not-on-shopee, %d match-review locked."
          % (summary["linked"],summary["not_selling"],summary["not_on_shopee"],summary["match_review_locked"]))
    s.cell(len(rows)+6,2,note).font=Font(name="Arial",italic=True,size=10,color="808080")
    for col,w in {"A":2,"B":40,"C":12}.items(): s.column_dimensions[col].width=w

    def sheet(title, headers, kinds, rows_data, widths, edit_cols=(), dv_col=None, dv_list=None):
        ws=wb.create_sheet(title); _style(ws,headers,kinds)
        for i,row in enumerate(rows_data,start=2):
            for j,val in enumerate(row,start=1):
                cell=ws.cell(i,j,val); cell.font=base
                if j in edit_cols: cell.fill=editf
        for i,w in enumerate(widths,start=1): ws.column_dimensions[get_column_letter(i)].width=w
        ws.freeze_panes="A2"
        if ws.max_row>=2: ws.auto_filter.ref=f"A1:{get_column_letter(len(headers))}{ws.max_row}"
        if dv_col and dv_list and ws.max_row>=2:
            dv=DataValidation(type="list",formula1='"%s"'%",".join(dv_list),allow_blank=True)
            ws.add_data_validation(dv); dv.add(f"{get_column_letter(dv_col)}2:{get_column_letter(dv_col)}{ws.max_row}")
        return ws

    # Locked Matches
    sheet("Locked Matches",
        ["#","Shopee Product ID","Shopee Product Name","Variation ID","Variation Name","SKU",
         "LOCKED Masterlist ID(s)","ML Model(s)|Color","ML Available Qty","Target Stock","# SKUs"],
        ["N","S","S","S","S","S","M","M","M","M","N"],
        [[i,x["pid"],x["pname"],x["vid"],x["vname"],x["sku"],"; ".join(x["mids"]),
          x["ml_desc"],"; ".join(str(mlmap.get(m,{}).get("avail","?")) for m in x["mids"]),
          x["target"],len(x["mids"])] for i,x in enumerate(cls["locked"],1)],
        [5,16,30,13,26,10,20,34,14,11,7])
    # New Masterlist SKUs (editable)
    sheet("New Masterlist SKUs",
        ["#","Masterlist Stock Type ID","Category","Brand","Model","Color","Available Qty",
         "Link to Shopee Variation ID","Reviewer Decision","Notes"],
        ["N","M","M","M","M","M","M","S","N","N"],
        [[i,m["id"],m["category"],m["brand"],m["model"],m["color"],m["avail"],"","",""] for i,m in enumerate(new_ml,1)],
        [5,22,10,12,30,20,12,26,22,22], edit_cols=(8,9,10), dv_col=9,
        dv_list=["Linked (fill col H)","Not applicable","Not on Shopee yet"])
    # Match Review (editable) — Shopee SKUs needing a match (incl. brand-new ones)
    sheet("Match Review",
        ["#","Shopee Product ID","Shopee Product Name","Variation ID","Variation Name","SKU","Current Seller Stock",
         "Corrected Masterlist ID","Reviewer Decision","Notes"],
        ["N","S","S","S","S","S","S","M","N","N"],
        [[i,x["pid"],x["pname"],x["vid"],x["vname"],x["sku"],x["stock"],"","",""] for i,x in enumerate(cls["match_review"],1)],
        [5,16,32,13,28,10,10,22,22,22], edit_cols=(8,9,10), dv_col=9,
        dv_list=["Linked (fill col H)","Not on Masterlist"])
    # Skipped
    sheet("Skipped (SKU=NO or 0 Qty)",
        ["#","Shopee Product ID","Shopee Product Name","Variation ID","Variation Name","SKU","Target Stock"],
        ["N","S","S","S","S","S","N"],
        [[i,x["pid"],x["pname"],x["vid"],x["vname"],x["sku"],0] for i,x in enumerate(cls["skipped"],1)],
        [5,16,40,13,30,12,10])
    # Accessories
    sheet("Accessories (not matched)",
        ["#","Shopee Product ID","Shopee Product Name","Variation ID","Variation Name","SKU","Target Stock (preset)"],
        ["N","S","S","S","S","S","N"],
        [[i,x["pid"],x["pname"],x["vid"],x["vname"],x["sku"],ACC_PRESET] for i,x in enumerate(cls["accessories"],1)],
        [5,16,44,13,30,12,16])
    # Not Selling / Not on Shopee Yet
    for title,ids in [("Not Selling in Shopee",not_selling),("Not on Shopee Yet",not_on_shopee)]:
        sheet(title,["#","Masterlist Stock Type ID","Category","Brand","Model","Color","Available Qty"],
            ["N","M","M","M","M","M","M"],
            [[i,mid,mlmap.get(mid,{}).get("category",""),mlmap.get(mid,{}).get("brand",""),
              mlmap.get(mid,{}).get("model",""),mlmap.get(mid,{}).get("color",""),mlmap.get(mid,{}).get("avail","")]
             for i,mid in enumerate(sorted(ids),1)],
            [5,22,10,12,30,20,12])
    out=io.BytesIO(); wb.save(out); out.seek(0); return out.getvalue()

# ---------- Write updated Shopee file ----------
def process_shopee(data, avail_map, locked, file_label="file"):
    wb=openpyxl.load_workbook(fix_shopee_bytes(data)); ws=wb.active
    vcol=scol=skuc=pnc=hr=None
    for r in range(1,min(4,ws.max_row)+1):
        vals=[str(c.value).strip().lower() if c.value is not None else "" for c in ws[r]]
        for i,v in enumerate(vals):
            if v=="et_title_variation_id" or v=="variation id": vcol=i+1
            if v.startswith("et_title_variation_sto") or v=="seller stock": scol=i+1
            if v=="et_title_variation_sku" or v=="sku": skuc=i+1
            if v=="et_title_product_name" or v=="product name": pnc=i+1
        if vcol and scol: hr=r; break
    if not (vcol and scol): raise ValueError(f"Shopee {file_label}: missing Variation ID / Seller Stock.")
    stats={"total_rows":0,"locked":0,"skipped":0,"accessories":0,"unchanged":0}
    for r in range(hr+1,ws.max_row+1):
        vid=_nid(ws.cell(r,vcol).value)
        if not vid or not vid.isdigit(): continue
        stats["total_rows"]+=1; cell=ws.cell(r,scol)
        sku=ws.cell(r,skuc).value if skuc else ""; pn=ws.cell(r,pnc).value if pnc else ""
        if vid in locked:
            cell.value=sum(avail_map.get(m,0) for m in locked[vid]); stats["locked"]+=1
        elif is_skip_sku(sku):
            cell.value=0; stats["skipped"]+=1
        elif is_accessory(pn):
            cell.value=ACC_PRESET; stats["accessories"]+=1
        else:
            stats["unchanged"]+=1
    out=io.BytesIO(); wb.save(out); out.seek(0)
    return out.getvalue(), stats
